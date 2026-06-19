"""Doc file attribute (ten doan -> tiet dien thep to hop).

Chua co dinh dang co dinh nen ho tro vai dinh dang pho bien:
  - .csv / .txt: 2 cot, cot 1 la ten doan, cot 2 la tiet dien (phan cach bang dau , hoac ;)
  - .xlsx: 2 cot dau tien cua sheet dau tien

Goi load_attributes(path) -> dict ten_doan -> tiet_dien
"""
from __future__ import annotations

import csv
from pathlib import Path


def load_attributes(path: str) -> dict[str, str]:
    p = Path(path)
    if p.suffix.lower() in (".csv", ".txt"):
        return _load_csv(p)
    if p.suffix.lower() in (".xlsx", ".xls"):
        return _load_excel(p)
    raise ValueError(f"Chua ho tro dinh dang file attribute: {p.suffix}")


def _load_csv(p: Path) -> dict[str, str]:
    mapping = {}
    with open(p, newline="", encoding="utf-8-sig") as f:
        sample = f.read(2048)
        f.seek(0)
        delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.reader(f, delimiter=delimiter)
        for row in reader:
            if len(row) < 2:
                continue
            name, section = row[0].strip(), row[1].strip()
            if not name or name.lower() in ("name", "ten", "ten doan"):
                continue
            mapping[name] = section
    return mapping


def _load_excel(p: Path) -> dict[str, str]:
    from openpyxl import load_workbook

    wb = load_workbook(p, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    mapping = {}
    for row in ws.iter_rows(values_only=True):
        if not row or row[0] is None:
            continue
        name = str(row[0]).strip()
        section = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
        if not name or name.lower() in ("name", "ten", "ten doan"):
            continue
        mapping[name] = section
    return mapping
